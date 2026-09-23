import io
import re
import zipfile
from dataclasses import dataclass
from pathlib import Path
from uuid import uuid4

from docx import Document
from fastapi import HTTPException, UploadFile, status
from pypdf import PdfReader

from app.core.config import Settings

ALLOWED_EXTENSIONS = {"pdf", "docx", "txt"}
CANONICAL_MIME_TYPES = {
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "txt": "text/plain",
}
ACCEPTED_MIME_TYPES = {
    "pdf": {"application/pdf", "application/octet-stream"},
    "docx": {
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "application/zip",
        "application/octet-stream",
    },
    "txt": {"text/plain", "application/octet-stream"},
}
MAX_DOCX_ENTRIES = 2_000
MAX_DOCX_UNCOMPRESSED_BYTES = 50 * 1024 * 1024
SIGNALS = (
    "moet",
    "dient",
    "verplicht",
    "vereist",
    "inschrijver toont aan",
    "opdrachtnemer waarborgt",
    "leverancier beschrijft",
    "bewijsstuk",
    "certificaat",
    "beveiliging",
    "persoonsgegevens",
    "algoritme",
    "ai-systeem",
    "continuïteit",
    "subverwerker",
)


@dataclass(frozen=True)
class SavedUpload:
    original_name: str
    stored_name: str
    mime_type: str
    extracted_text: str


@dataclass(frozen=True)
class RequirementCandidate:
    text: str
    title: str
    category: str
    location: str
    fragment: str
    priority: str


def _safe_name(filename: str | None) -> tuple[str, str]:
    basename = Path(filename or "").name
    sanitized = re.sub(r"[^a-zA-Z0-9._ -]", "_", basename).strip(" .")
    extension = sanitized.rsplit(".", 1)[-1].lower() if "." in sanitized else ""
    if not sanitized or extension not in ALLOWED_EXTENSIONS:
        raise HTTPException(
            status_code=status.HTTP_415_UNSUPPORTED_MEDIA_TYPE,
            detail="Alleen PDF-, DOCX- en TXT-bestanden zijn toegestaan.",
        )
    return sanitized[:255], extension


def _extract_text(data: bytes, extension: str) -> str:
    try:
        if extension == "txt":
            try:
                return data.decode("utf-8-sig").replace("\x00", "")
            except UnicodeDecodeError:
                return data.decode("cp1252").replace("\x00", "")
        if extension == "docx":
            document = Document(io.BytesIO(data))
            parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
            for table in document.tables:
                parts.extend(
                    " | ".join(cell.text.strip() for cell in row.cells)
                    for row in table.rows
                    if any(cell.text.strip() for cell in row.cells)
                )
            return "\n\n".join(parts)
        reader = PdfReader(io.BytesIO(data))
        pages = []
        for index, page in enumerate(reader.pages, start=1):
            pages.append(f"[[Pagina {index}]]\n{page.extract_text() or ''}")
        return "\n\n".join(pages)
    except Exception as exc:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail="De tekst kon niet uit het bestand worden gelezen.",
        ) from exc


def _validate_content_type(extension: str, content_type: str | None) -> None:
    declared = (content_type or "").lower().split(";", 1)[0].strip()
    if declared and declared not in ACCEPTED_MIME_TYPES[extension]:
        raise HTTPException(
            status_code=status.HTTP_415_UNSUPPORTED_MEDIA_TYPE,
            detail="Het opgegeven bestandstype komt niet overeen met de bestandsextensie.",
        )


def _validate_txt(data: bytes) -> None:
    known_binary_headers = (b"%PDF-", b"PK\x03\x04", b"\x89PNG", b"\xff\xd8\xff", b"GIF8")
    if data.startswith(known_binary_headers) or b"\x00" in data:
        raise HTTPException(
            status_code=status.HTTP_415_UNSUPPORTED_MEDIA_TYPE,
            detail="Het bestand heeft geen geldige TXT-inhoud.",
        )
    control_bytes = sum(byte < 32 and byte not in {9, 10, 13} for byte in data)
    if control_bytes / len(data) > 0.02:
        raise HTTPException(
            status_code=status.HTTP_415_UNSUPPORTED_MEDIA_TYPE,
            detail="Het bestand lijkt binaire inhoud te bevatten en is geen geldig TXT-bestand.",
        )


def _validate_docx(data: bytes) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            entries = archive.infolist()
            names = {entry.filename for entry in entries}
            if "[Content_Types].xml" not in names or "word/document.xml" not in names:
                raise ValueError("Vereiste DOCX-onderdelen ontbreken.")
            if len(entries) > MAX_DOCX_ENTRIES:
                raise ValueError("Het DOCX-bestand bevat te veel onderdelen.")
            if sum(entry.file_size for entry in entries) > MAX_DOCX_UNCOMPRESSED_BYTES:
                raise ValueError("Het uitgepakte DOCX-bestand is te groot.")
            for entry in entries:
                path = Path(entry.filename)
                if path.is_absolute() or ".." in path.parts:
                    raise ValueError("Het DOCX-bestand bevat een onveilig pad.")
            if archive.testzip() is not None:
                raise ValueError("Het DOCX-archief is beschadigd.")
    except (ValueError, zipfile.BadZipFile, OSError) as exc:
        raise HTTPException(
            status_code=status.HTTP_415_UNSUPPORTED_MEDIA_TYPE,
            detail="Het bestand heeft geen geldige DOCX-structuur.",
        ) from exc


def _validate_content(data: bytes, extension: str, content_type: str | None) -> str:
    _validate_content_type(extension, content_type)
    if extension == "pdf" and not data[:1024].lstrip().startswith(b"%PDF-"):
        raise HTTPException(
            status_code=status.HTTP_415_UNSUPPORTED_MEDIA_TYPE,
            detail="Het bestand heeft geen geldige PDF-header.",
        )
    if extension == "docx":
        _validate_docx(data)
    if extension == "txt":
        _validate_txt(data)
    return CANONICAL_MIME_TYPES[extension]


async def save_upload(project_id: str, upload: UploadFile, settings: Settings) -> SavedUpload:
    original_name, extension = _safe_name(upload.filename)
    data = await upload.read(settings.max_upload_bytes + 1)
    if not data:
        raise HTTPException(status_code=422, detail="Kies een bestand met inhoud.")
    if len(data) > settings.max_upload_bytes:
        raise HTTPException(status_code=413, detail="Het bestand is groter dan 10 MB.")
    verified_mime_type = _validate_content(data, extension, upload.content_type)
    extracted_text = _extract_text(data, extension)

    safe_project_id = re.sub(r"[^a-zA-Z0-9_-]", "", project_id)
    if safe_project_id != project_id:
        raise HTTPException(status_code=400, detail="Ongeldige projectreferentie.")

    stored_name = f"{uuid4()}.{extension}"
    directory = Path(settings.upload_dir).resolve() / safe_project_id
    directory.mkdir(parents=True, exist_ok=True)
    target = (directory / stored_name).resolve()
    if directory not in target.parents:
        raise HTTPException(status_code=400, detail="Ongeldige bestandsnaam.")
    with target.open("xb") as stored_file:
        stored_file.write(data)
    return SavedUpload(
        original_name=original_name,
        stored_name=stored_name,
        mime_type=verified_mime_type,
        extracted_text=extracted_text,
    )


def discard_upload(project_id: str, upload: SavedUpload, settings: Settings) -> None:
    safe_project_id = re.sub(r"[^a-zA-Z0-9_-]", "", project_id)
    if safe_project_id != project_id:
        return
    upload_root = Path(settings.upload_dir).resolve()
    directory = (upload_root / safe_project_id).resolve()
    target = (directory / upload.stored_name).resolve()
    if directory not in target.parents:
        return
    target.unlink(missing_ok=True)
    try:
        directory.rmdir()
    except OSError:
        pass


def categorize(text: str) -> str:
    lowered = text.lower()
    if re.search(r"encrypt|versleutel|toegang|logging|beveilig|pentest", lowered):
        return "informatiebeveiliging"
    if re.search(r"persoonsgegeven|bewaartermijn|subverwerk|\beer\b|privacy", lowered):
        return "privacy"
    if re.search(r"algorit|ai-systeem|model|menselijke tussenkomst|uitleg", lowered):
        return "AI en algoritmen"
    if re.search(r"continuïteit|herstel|beschikbaarheid", lowered):
        return "continuïteit"
    if re.search(r"governance|organisatie|verantwoordelijk", lowered):
        return "organisatie en governance"
    if re.search(r"contract|aansprak|voorwaarde", lowered):
        return "juridische voorwaarden"
    if re.search(r"duurzaam|milieu|energie", lowered):
        return "duurzaamheid"
    if re.search(r"financ|omzet|verzekering", lowered):
        return "financieel"
    return "overig"


def short_title(text: str) -> str:
    title = " ".join(text.rstrip(".:;").split()[:8])
    return f"{title[:59]}…" if len(title) > 62 else title


def extract_requirements(text: str) -> list[RequirementCandidate]:
    candidates: list[RequirementCandidate] = []
    seen: set[str] = set()
    current_page: str | None = None
    paragraph_number = 0
    for block in re.split(r"\n\s*\n|\n", text):
        block = block.strip()
        page = re.fullmatch(r"\[\[Pagina (\d+)]]", block)
        if page:
            current_page = f"Pagina {page.group(1)}"
            continue
        if not block:
            continue
        paragraph_number += 1
        location = current_page or f"Alinea {paragraph_number}"
        for sentence in re.split(r"(?<=[.!?])\s+", block):
            clean = re.sub(r"^[-•\d.)\s]+", "", sentence).strip()
            lowered = clean.lower()
            if len(clean) <= 20 or not any(signal in lowered for signal in SIGNALS):
                continue
            key = clean.casefold()
            if key in seen:
                continue
            seen.add(key)
            candidates.append(
                RequirementCandidate(
                    text=clean,
                    title=short_title(clean),
                    category=categorize(clean),
                    location=location,
                    fragment=clean,
                    priority="wens"
                    if re.search(r"\bwens\b|bij voorkeur", lowered)
                    else "verplicht",
                )
            )
            if len(candidates) == 80:
                return candidates
    return candidates
